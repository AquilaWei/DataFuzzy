"""Writing charts into the school's Excel form: one workbook per student and location, a
sheet per school year and IEP or 504 plan, each visit a row of the template: date, time,
session type, activity, English chart and a score for each goal of the plan."""

from __future__ import annotations

import math
import re
import unicodedata
from collections.abc import Iterable
from copy import copy
from datetime import date
from importlib.resources import as_file, files
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .db import Plan, Record, RecordStore, Student, StudentProfile, parse_plan_date
from .school_year import default_school_year, format_school_year, parse_school_year

TEMPLATE_SHEET = "blank"
HEADER_ROW, GOALS_ROW, HEADINGS_ROW = 1, 2, 3
FIRST_SLOT_ROW = 4  # one row per visit
TEMPLATE_SLOTS = 8
# Columns of a visit's row; Session type, Activity and Comment span several (see _merge_slot).
DATE_COL, TIME_COL, TYPE_COL, ACTIVITY_COL, COMMENT_COL, FIRST_GOAL_COL = 1, 2, 3, 5, 7, 13
TEMPLATE_GOALS = 2  # the template's goal columns, M and N
BLANK_COL = 15  # a column of the template with nothing in it, for resetting spare goal columns
DATE_FORMAT = "yyyy-mm-dd"
LINE_HEIGHT = 16  # points per line of 12-point text
NO_SCORE = "--"

_FILE_UNSAFE = re.compile(r'[/\\:*?"<>|\x00-\x1f]')
_SHEET_UNSAFE = re.compile(r"[\[\]:*?/\\]")
SHEET_TITLE_MAX = 31


def group_by_location(records: Iterable[Record]) -> tuple[dict[str, list[Record]], list[Record]]:
    """Records per location, in the order given, and the records that have no location."""
    groups: dict[str, list[Record]] = {}
    unassigned: list[Record] = []
    for record in records:
        if record.location.strip():
            groups.setdefault(record.location.strip(), []).append(record)
        else:
            unassigned.append(record)
    return groups, unassigned


def _safe(text: str, fallback: str) -> str:
    return _FILE_UNSAFE.sub("", text).strip(" .") or fallback


def file_name(name: str, date_from: date, date_to: date, location: str | None = None) -> str:
    """`<name>-<from>至<to>.xlsx`, or `<location>-<name>-…` to tell apart two files of one name,
    with characters no file system accepts removed."""
    prefix = f"{_safe(location, 'location')}-" if location is not None else ""
    return f"{prefix}{_safe(name, 'student')}-{date_from.isoformat()}至{date_to.isoformat()}.xlsx"


def sheet_title(name: str, used: set[str]) -> str:
    """A valid, unused sheet title for `name`; adds it to `used` (compared case-insensitively,
    as Excel does)."""
    base = _SHEET_UNSAFE.sub("", name).strip("' ") or "student"
    title = base[:SHEET_TITLE_MAX]
    n = 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[:SHEET_TITLE_MAX - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def school_year_of(day: date, profile: StudentProfile) -> tuple[date, date]:
    """The school year `day` falls in, its months those of the profile's school year (shifted
    by whole years) or, when not set, September to June starting from August."""
    months = parse_school_year(profile.school_year)
    if months is None:
        return default_school_year(day)
    start, end = months
    first = day.year if day.month >= start.month else day.year - 1
    shift = first - start.year
    return date(start.year + shift, start.month, 1), date(end.year + shift, end.month, 1)


def period_title(school_year: tuple[date, date], plan: Plan) -> str:
    """"2026-2027_IEP 10∕29∕26": the school year and the plan of a sheet. Excel takes no "/"
    in a sheet's name, so the date's slashes are "∕" (division slash), which looks the same."""
    title = f"{school_year[0].year}-{school_year[1].year}_{plan.kind}"
    day = parse_plan_date(plan.plan_date)
    return f"{title} {day:%m}∕{day:%d}∕{day:%y}" if day else title


def periods(records: Iterable[Record], profile: StudentProfile) -> list[tuple[tuple[date, date], Plan, list[Record]]]:
    """A student's records split where the school year or the plan in effect changes, oldest
    first: (school year, plan, its records oldest first, a day's by start time) each."""
    groups: dict[tuple[tuple[date, date], Plan], list[Record]] = {}
    for record in sorted(records, key=lambda r: (r.visit_date, r.start_time or "99:99", r.id)):
        key = (school_year_of(record.visit_date, profile), profile.plan_on(record.visit_date))
        groups.setdefault(key, []).append(record)
    return [(year, plan, rs) for (year, plan), rs in groups.items()]


def header_text(name: str, profile: StudentProfile, today: date, plan: Plan | None = None,
                school_year: tuple[date, date] | None = None) -> str:
    """The form's first row: who the student is, the date of their IEP (or 504 plan; the
    latest unless `plan` is given) and the school year (`school_year`, else the profile's,
    else this one) and the plan's frequency ("30 mins/month"; blank when not set)."""
    plan = plan or profile.plan
    if school_year is not None:
        year_text = format_school_year(*school_year)
    else:
        year_text = profile.school_year or format_school_year(*default_school_year(today))
    return (f"Name: {name} / Diagnosis: {profile.diagnosis} / Age: {profile.age} / "
            f"Grade: {profile.grade} / Classroom: {profile.classroom} / "
            f"Case manager: {profile.case_manager}\n{plan.label}: {plan.plan_date}\n"
            f"Frequency: {plan.frequency_text}".rstrip() + f"\nSchool year: {year_text}")


def goals_text(goals: tuple[str, ...]) -> str:
    """The form's second row, "Goal 1: …" a line each; "Goal 1:" for a student with none."""
    return "\n".join(f"Goal {i}: {goal}" for i, goal in enumerate(goals or ("",), start=1)).rstrip()


def time_text(start: str, end: str) -> str:
    """"9:30-10:00" from "09:30" and "10:00", on the 12-hour clock as the form is written
    by hand; just the start without an end, "" without a start."""
    def short(hhmm: str) -> str:
        hour, minute = (int(part) for part in hhmm.split(":"))
        return f"{(hour - 1) % 12 + 1}:{minute:02d}"
    if not start:
        return ""
    return f"{short(start)}-{short(end)}" if end else short(start)


def score_text(score: str) -> str:
    """A goal's score as the form shows it: "75" becomes "75%", a blank one is NO_SCORE and
    anything else ("75%", "not met") stays as typed."""
    score = score.strip()
    if not score:
        return NO_SCORE
    return f"{score}%" if re.fullmatch(r"\d+(\.\d+)?", score) else score


def _text_width(text: str) -> int:
    """Width in Excel's character units: a CJK character takes two."""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _lines(text: str, width: float) -> int:
    """About how many lines `text` wraps to in a cell `width` characters wide."""
    usable = max(width * 0.9, 1)  # 12-point text is a little wider than the width unit
    return sum(max(1, math.ceil(_text_width(line) / usable)) for line in text.split("\n"))


def _width(ws: Worksheet, first: int, last: int) -> float:
    return sum(ws.column_dimensions[get_column_letter(c)].width or 13 for c in range(first, last + 1))


def _set_goal_columns(ws: Worksheet, count: int) -> int:
    """Make the template's two goal columns `count` (at least one), each styled like the
    template's, and stretch the two header rows over them. Returns the last column."""
    last = FIRST_GOAL_COL + max(count, 1) - 1
    template_last = FIRST_GOAL_COL + TEMPLATE_GOALS - 1
    for cells in (f"A{HEADER_ROW}:{get_column_letter(template_last)}{HEADER_ROW}",
                  f"A{GOALS_ROW}:{get_column_letter(template_last - 1)}{GOALS_ROW}"):
        ws.unmerge_cells(cells)
    rows = range(HEADER_ROW, FIRST_SLOT_ROW + TEMPLATE_SLOTS)
    # Taken before any is changed, as the template's own goal columns are restyled too.
    styles = {(row, col): copy(ws.cell(row, col)._style)
              for row in rows for col in (2, FIRST_GOAL_COL, template_last, BLANK_COL)}
    width = ws.column_dimensions[get_column_letter(FIRST_GOAL_COL)].width or 13
    for col in range(FIRST_GOAL_COL, max(last, template_last) + 1):
        for row in rows:
            if col > last:
                source = (row, BLANK_COL)
            elif row in (HEADER_ROW, GOALS_ROW):  # both like the header: one framed box across
                source = (HEADER_ROW, template_last if col == last else FIRST_GOAL_COL)
            else:
                source = (row, FIRST_GOAL_COL)
            ws.cell(row, col)._style = copy(styles[source])
        ws.column_dimensions[get_column_letter(col)].width = width
        ws.cell(HEADINGS_ROW, col).value = f"Goal {col - FIRST_GOAL_COL + 1}" if col <= last else None
    for col in range(2, min(last, FIRST_GOAL_COL)):
        ws.cell(GOALS_ROW, col)._style = copy(styles[(HEADER_ROW, 2)])
    for row in (HEADER_ROW, GOALS_ROW):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    return last


def _fill_header(ws: Worksheet, student: Student, profile: StudentProfile, today: date,
                 plan: Plan, school_year: tuple[date, date]) -> int:
    """Fill the two header rows for `plan` in `school_year`; returns the last column of the form."""
    last = _set_goal_columns(ws, len(plan.goals))
    ws.cell(HEADER_ROW, 1).value = header_text(student.name, profile, today, plan, school_year)
    goals = ws.cell(GOALS_ROW, 1)
    goals.value = goals_text(plan.goals)
    goals.alignment = Alignment(vertical="top", wrap_text=True)
    lines = _lines(goals.value, _width(ws, 1, last))
    ws.row_dimensions[GOALS_ROW].height = max(ws.row_dimensions[GOALS_ROW].height or 0, LINE_HEIGHT * lines)
    return last


def _load_template() -> Workbook:
    with as_file(files("datafuzzy.ot").joinpath("excel_template.xlsx")) as path:
        return load_workbook(path)


def _copy_slot_style(ws: Worksheet, slot: int, last: int) -> None:
    """Give slot `slot` (0-based, past the template's) the look of the last template slot."""
    source = FIRST_SLOT_ROW + TEMPLATE_SLOTS - 1
    for col in range(1, last + 1):
        ws.cell(FIRST_SLOT_ROW + slot, col)._style = copy(ws.cell(source, col)._style)
    _merge_slot(ws, FIRST_SLOT_ROW + slot)


def _merge_slot(ws: Worksheet, row: int) -> None:
    for cells in (f"C{row}:D{row}", f"E{row}:F{row}", f"G{row}:L{row}"):
        if cells not in ws.merged_cells:
            ws.merge_cells(cells)


def _fill(ws: Worksheet, records: list[Record], goal_count: int, last: int) -> None:
    spans = {DATE_COL: DATE_COL, TIME_COL: TIME_COL, TYPE_COL: ACTIVITY_COL - 1,
             ACTIVITY_COL: COMMENT_COL - 1, COMMENT_COL: FIRST_GOAL_COL - 1}
    for slot, record in enumerate(records):
        if slot >= TEMPLATE_SLOTS:
            _copy_slot_style(ws, slot, last)
        row = FIRST_SLOT_ROW + slot
        scores = [score_text(s) for s in record.goal_scores[:goal_count]]
        scores += [NO_SCORE] * (max(goal_count, 1) - len(scores))
        cells = [(DATE_COL, record.visit_date), (TIME_COL, time_text(record.start_time, record.end_time)),
                 (TYPE_COL, record.session_type), (ACTIVITY_COL, record.activity),
                 (COMMENT_COL, record.en_text)]
        cells += [(FIRST_GOAL_COL + i, score) for i, score in enumerate(scores)]
        lines = 1
        for col, value in cells:
            cell = ws.cell(row, col)
            cell.value = value if value != "" else None
            cell.alignment = Alignment(horizontal=cell.alignment.horizontal,
                                       vertical="top", wrap_text=True)
            if isinstance(value, str) and value:
                lines = max(lines, _lines(value, _width(ws, col, spans.get(col, col))))
        ws.cell(row, DATE_COL).number_format = DATE_FORMAT
        if lines > 1:  # Excel does not fit a row's height to merged cells by itself
            ws.row_dimensions[row].height = LINE_HEIGHT * lines


def build_workbook(store: RecordStore, records: Iterable[Record], today: date | None = None) -> Workbook:
    """A sheet per student, school year and plan: a new sheet where the school year changes
    or a new IEP or 504 plan begins, named like "2026-2027_IEP 10∕29∕26" (after the student's
    name when there are several students). Visits oldest first, the header and goals those of
    the sheet's plan; the template sheet is dropped."""
    by_student: dict[int, list[Record]] = {}
    for record in records:
        by_student.setdefault(record.student_id, []).append(record)
    students = sorted((store.get_student(sid) for sid in by_student), key=lambda s: (s.name, s.id))
    wb = _load_template()
    template = wb[TEMPLATE_SHEET]
    used: set[str] = set()
    for student in students:
        profile = store.get_profile(student.id)
        for school_year, plan, student_records in periods(by_student[student.id], profile):
            ws = wb.copy_worksheet(template)
            title = period_title(school_year, plan)
            ws.title = sheet_title(title if len(students) == 1 else f"{student.name} {title}", used)
            # copy_worksheet leaves out the print settings.
            ws.page_setup.orientation = template.page_setup.orientation
            ws.page_margins = copy(template.page_margins)
            ws.sheet_view.zoomScale = template.sheet_view.zoomScale
            last = _fill_header(ws, student, profile, today or date.today(), plan, school_year)
            _fill(ws, student_records, len(plan.goals), last)
    wb.remove(template)
    for i, ws in enumerate(wb.worksheets):
        ws.sheet_view.tabSelected = i == 0
    if wb.worksheets:
        wb.active = 0
    return wb


def write_workbook(store: RecordStore, records: Iterable[Record], path: Path) -> None:
    build_workbook(store, records).save(path)
