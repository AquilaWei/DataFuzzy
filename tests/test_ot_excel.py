from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QDate

from datafuzzy.ot.db import Plan, RecordStore, StudentProfile
from datafuzzy.ot.excel_export import (
    build_workbook,
    file_name,
    goals_text,
    group_by_location,
    header_text,
    period_title,
    school_year_of,
    score_text,
    sheet_title,
    time_text,
    write_workbook,
)
from datafuzzy.ot.ui.exchange_dialogs import ExcelExportDialog

KEY = bytes(range(32))
TODAY = date(2026, 10, 25)


@pytest.fixture
def store(tmp_path):
    with RecordStore(tmp_path / "t.db", KEY) as s:
        yield s


@pytest.fixture
def charts(store):
    ming = store.add_student("王小明")
    amy = store.add_student("Amy")
    store.add_record(ming.id, date(2026, 10, 8), "Lincoln", "中文", "Worked on cutting.", activity="剪紙")
    store.add_record(ming.id, date(2026, 10, 1), "Lincoln", "中文", "Drew circles.", activity="畫畫")
    store.add_record(amy.id, date(2026, 10, 2), "Lincoln", activity="鞦韆")
    store.add_record(amy.id, date(2026, 10, 3), "新竹國小", en_text="Balance.")
    store.add_record(amy.id, date(2026, 10, 4), "", en_text="No place.")
    return store


def test_file_name_is_the_student_and_the_exported_period():
    assert file_name("王小明", date(2026, 10, 1), date(2026, 10, 31)) == "王小明-2026-10-01至2026-10-31.xlsx"


def test_file_name_drops_characters_no_file_system_accepts():
    assert file_name("Amy/Lee: ?", date(2026, 10, 1), date(2026, 10, 31)) == "AmyLee-2026-10-01至2026-10-31.xlsx"


def test_file_name_can_lead_with_the_location():
    assert (file_name("Amy", date(2026, 10, 1), date(2026, 10, 31), "Lincoln")
            == "Lincoln-Amy-2026-10-01至2026-10-31.xlsx")


def test_sheet_titles_are_valid_and_unique():
    used: set[str] = set()
    assert sheet_title("Amy/Lee", used) == "AmyLee"
    assert sheet_title("amylee", used) == "amylee (2)"
    long = sheet_title("x" * 40, used)
    assert len(long) == 31
    assert sheet_title("x" * 40, used) == "x" * 27 + " (2)"


def test_records_are_grouped_by_location(charts):
    groups, unassigned = group_by_location(charts.records_between(date(2026, 10, 1), date(2026, 10, 31)))
    assert {k: len(v) for k, v in groups.items()} == {"Lincoln": 3, "新竹國小": 1}
    assert [r.en_text for r in unassigned] == ["No place."]


def test_workbook_has_a_sheet_per_student_filled_oldest_first(charts, tmp_path):
    groups, _ = group_by_location(charts.records_between(date(2026, 10, 1), date(2026, 10, 31)))
    path = tmp_path / "out.xlsx"
    write_workbook(charts, groups["Lincoln"], path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Amy 2026-2027_IEP", "王小明 2026-2027_IEP"]
    ws = wb["王小明 2026-2027_IEP"]
    assert ws["A1"].value.startswith("Name:")
    assert ws["A2"].value == "Goal 1:"
    assert (ws["A4"].value, ws["E4"].value, ws["G4"].value) == (datetime(2026, 10, 1), "畫畫", "Drew circles.")
    assert (ws["A5"].value, ws["E5"].value, ws["G5"].value) == (datetime(2026, 10, 8), "剪紙", "Worked on cutting.")
    assert ws["A6"].value is None
    assert {"C4:D4", "E4:F4", "G4:L4"} <= {str(r) for r in ws.merged_cells.ranges}
    assert (wb["Amy 2026-2027_IEP"]["E4"].value, wb["Amy 2026-2027_IEP"]["G4"].value) == ("鞦韆", None)


def test_header_lists_the_profile_iep_date_and_school_year():
    profile = StudentProfile(diagnosis="ADHD, ODD", age="8", grade="3", classroom="301",
                             case_manager="Ms. Lee", school_year="09/2026-06/2027", plan_date="10/15/2026")
    assert header_text("王小明", profile, TODAY) == (
        "Name: 王小明 / Diagnosis: ADHD, ODD / Age: 8 / Grade: 3 / Classroom: 301 / "
        "Case manager: Ms. Lee\nIEP Date: 10/15/2026\nFrequency:\nSchool year: 09/2026-06/2027")


def test_header_of_a_504_plan_shows_the_504_date():
    profile = StudentProfile(plan_504=True, plan_date="03/02/2026")
    assert "\n504 Date: 03/02/2026\n" in header_text("Amy", profile, TODAY)


def test_goals_row_has_a_line_per_goal():
    assert goals_text(("Fine motor", "Self-care")) == "Goal 1: Fine motor\nGoal 2: Self-care"


def test_time_is_written_as_the_form_has_it():
    assert time_text("09:30", "10:00") == "9:30-10:00"


def test_afternoon_time_is_on_the_12_hour_clock():
    assert time_text("13:00", "13:30") == "1:00-1:30"


def test_time_without_an_end_is_the_start():
    assert time_text("09:30", "") == "9:30"


def test_a_bare_number_score_is_a_percentage():
    assert score_text("75") == "75%"


def test_a_blank_score_is_two_dashes():
    assert score_text(" ") == "--"


def test_not_met_stays_as_typed():
    assert score_text("not met") == "not met"


def test_header_without_a_school_year_uses_the_current_one():
    assert header_text("Amy", StudentProfile(), date(2027, 3, 1)).endswith("School year: 09/2026-06/2027")


def test_sheet_header_is_filled_from_the_students_profile(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile(diagnosis="ADHD", goals=("Fine motor", "Self-care")))
    store.add_record(student.id, date(2026, 10, 1), "Lincoln", en_text="Drew.")
    ws = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets[0]
    assert ws["A1"].value.startswith("Name: 王小明 / Diagnosis: ADHD / ")
    assert ws["A2"].value == "Goal 1: Fine motor\nGoal 2: Self-care"


def test_visit_row_has_time_session_type_and_a_score_per_goal(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile(goals=("Cut", "Write", "Button")))
    store.add_record(student.id, date(2026, 10, 1), start_time="09:30", end_time="10:00",
                     session_type="group", goal_scores=("75", "not met"))
    ws = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets[0]
    assert [ws.cell(4, c).value for c in (2, 3, 13, 14, 15)] == ["9:30-10:00", "group", "75%", "not met", "--"]


def test_each_goal_gets_a_column(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile(goals=("Cut", "Write", "Button")))
    store.add_record(student.id, date(2026, 10, 1))
    ws = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets[0]
    assert [ws.cell(3, c).value for c in (13, 14, 15, 16)] == ["Goal 1", "Goal 2", "Goal 3", None]
    assert ws["O4"].border.right.style == "thin"
    assert {"A1:O1", "A2:O2"} <= {str(r) for r in ws.merged_cells.ranges}


def test_one_goal_leaves_one_goal_column(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile(goals=("Cut",)))
    store.add_record(student.id, date(2026, 10, 1))
    ws = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets[0]
    assert (ws["M3"].value, ws["N3"].value, ws["N4"].border.right.style) == ("Goal 1", None, None)
    assert "A1:M1" in {str(r) for r in ws.merged_cells.ranges}


def test_record_from_before_session_types_exports_as_individual(store):
    student = store.add_student("王小明")
    store.add_record(student.id, date(2026, 10, 1))
    ws = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets[0]
    assert (ws["C4"].value, ws["M4"].value) == ("individual", "--")


def test_a_long_comment_makes_its_row_taller(store):
    student = store.add_student("王小明")
    store.add_record(student.id, date(2026, 10, 1), en_text="word " * 60)
    ws = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets[0]
    assert ws.row_dimensions[4].height > 16


def test_more_visits_than_the_template_has_slots(store, tmp_path):
    student = store.add_student("王小明")
    for i in range(12):
        store.add_record(student.id, date(2026, 10, 1) + timedelta(days=i), "Lincoln", en_text=f"day {i}")
    path = tmp_path / "out.xlsx"
    write_workbook(store, store.records_for_student(student.id), path)
    ws = load_workbook(path).worksheets[0]
    assert ws["G14"].value == "day 10"
    assert ws["G15"].value == "day 11"
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"C15:D15", "E15:F15", "G15:L15"} <= merged
    assert ws["A15"].border.bottom.style == "thin"


def make_dialog(qtbot, store, confirm=lambda path: True, folder=""):
    dialog = ExcelExportDialog(store, today=TODAY, confirm=confirm, choose=lambda start: folder)
    qtbot.addWidget(dialog)
    dialog.from_edit.setDate(QDate(2026, 10, 1))
    dialog.to_edit.setDate(QDate(2026, 10, 31))
    return dialog


def pick(dialog, location, name):
    dialog.location_box.setCurrentText(location)
    dialog.student_box.setCurrentIndex(
        [dialog.student_box.itemText(i).split("（")[0] for i in range(dialog.student_box.count())].index(name))
    dialog.add_btn.click()


def table_rows(dialog):
    return [tuple(dialog.table.item(r, c).text() for c in range(3)) for r in range(dialog.table.rowCount())]


def test_location_drop_down_lists_locations_with_visits_in_the_period(qtbot, charts):
    dialog = make_dialog(qtbot, charts)
    assert dialog.locations() == ["Lincoln", "新竹國小"]
    assert "1 份病歷沒有填地點" in dialog.status.text()


def test_student_drop_down_lists_who_was_seen_at_the_location(qtbot, charts):
    dialog = make_dialog(qtbot, charts)
    dialog.location_box.setCurrentText("Lincoln")
    assert [dialog.student_box.itemText(i) for i in range(dialog.student_box.count())] == [
        "Amy（1 份）", "王小明（2 份）"]


def test_picked_students_are_added_to_the_list_once(qtbot, charts):
    dialog = make_dialog(qtbot, charts)
    pick(dialog, "Lincoln", "王小明")
    pick(dialog, "Lincoln", "王小明")
    pick(dialog, "新竹國小", "Amy")
    assert table_rows(dialog) == [("Lincoln", "王小明", "2"), ("新竹國小", "Amy", "1")]


def test_add_all_adds_everyone_seen_at_the_location(qtbot, charts):
    dialog = make_dialog(qtbot, charts)
    dialog.location_box.setCurrentText("Lincoln")
    dialog.add_all_btn.click()
    assert table_rows(dialog) == [("Lincoln", "Amy", "1"), ("Lincoln", "王小明", "2")]


def test_remove_takes_the_selected_row_off_the_list(qtbot, charts):
    dialog = make_dialog(qtbot, charts)
    pick(dialog, "Lincoln", "王小明")
    pick(dialog, "Lincoln", "Amy")
    dialog.table.selectRow(0)
    dialog.remove_btn.click()
    assert table_rows(dialog) == [("Lincoln", "Amy", "1")]


def test_export_needs_a_folder(qtbot, charts):
    dialog = make_dialog(qtbot, charts)
    pick(dialog, "Lincoln", "王小明")
    assert not dialog.export_btn.isEnabled()


def test_folder_is_remembered(qtbot, charts, tmp_path):
    dialog = make_dialog(qtbot, charts, folder=str(tmp_path))
    dialog.folder_btn.click()
    assert charts.excel_folder() == str(tmp_path)
    assert dialog.folder_label.text() == str(tmp_path)


def test_each_student_gets_a_file_with_only_that_locations_visits(qtbot, charts, tmp_path):
    charts.set_excel_folder(str(tmp_path))
    dialog = make_dialog(qtbot, charts)
    pick(dialog, "Lincoln", "王小明")
    pick(dialog, "新竹國小", "Amy")
    dialog.export_btn.click()
    ming = tmp_path / "王小明-2026-10-01至2026-10-31.xlsx"
    amy = tmp_path / "Amy-2026-10-01至2026-10-31.xlsx"
    assert dialog.written == [ming, amy]
    assert load_workbook(ming).sheetnames == ["2026-2027_IEP"]
    assert load_workbook(amy).worksheets[0]["G4"].value == "Balance."
    assert load_workbook(amy).worksheets[0]["A5"].value is None
    assert dialog.result() == dialog.DialogCode.Accepted


def test_one_student_at_two_locations_gets_the_location_in_front(qtbot, charts, tmp_path):
    charts.set_excel_folder(str(tmp_path))
    dialog = make_dialog(qtbot, charts)
    pick(dialog, "Lincoln", "Amy")
    pick(dialog, "新竹國小", "Amy")
    dialog.export_btn.click()
    assert [p.name for p in dialog.written] == [
        "Lincoln-Amy-2026-10-01至2026-10-31.xlsx", "新竹國小-Amy-2026-10-01至2026-10-31.xlsx"]


def test_dialog_asks_before_replacing_a_file(qtbot, charts, tmp_path):
    charts.set_excel_folder(str(tmp_path))
    old = tmp_path / "王小明-2026-10-01至2026-10-31.xlsx"
    old.write_bytes(b"old")
    asked = []
    dialog = make_dialog(qtbot, charts, confirm=lambda path: asked.append(path) or False)
    pick(dialog, "Lincoln", "王小明")
    dialog.export_btn.click()
    assert asked == [old]
    assert old.read_bytes() == b"old"
    assert dialog.written == []
    assert dialog.result() != dialog.DialogCode.Accepted


OLD_PLAN = Plan(False, "10/29/2025", ("Cut",))
NEW_PLAN = Plan(True, "10/29/2026", ("Button", "Write"))


def test_sheet_title_is_the_school_year_and_the_plan():
    assert period_title((date(2026, 9, 1), date(2027, 6, 1)), Plan(False, "10/29/2026")) == "2026-2027_IEP 10∕29∕26"


def test_sheet_title_of_a_504_plan():
    assert period_title((date(2027, 9, 1), date(2028, 6, 1)), Plan(True, "10/29/2027")) == "2027-2028_504 10∕29∕27"


def test_sheet_title_of_a_plan_without_a_date_is_the_school_year_and_kind():
    assert period_title((date(2026, 9, 1), date(2027, 6, 1)), Plan()) == "2026-2027_IEP"


def test_school_year_follows_the_profiles_months():
    profile = StudentProfile(school_year="08/2026-07/2027")
    assert school_year_of(date(2027, 8, 3), profile) == (date(2027, 8, 1), date(2028, 7, 1))


def test_school_year_without_one_set_starts_in_august():
    assert school_year_of(date(2027, 7, 31), StudentProfile()) == (date(2026, 9, 1), date(2027, 6, 1))


def test_a_new_plan_in_the_period_starts_a_new_sheet(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile(school_year="09/2026-06/2027").with_plans([OLD_PLAN, NEW_PLAN]))
    store.add_record(student.id, date(2026, 10, 28), en_text="Before.")
    store.add_record(student.id, date(2026, 10, 29), en_text="After.")
    wb = build_workbook(store, store.records_for_student(student.id), today=TODAY)
    assert wb.sheetnames == ["2026-2027_IEP 10∕29∕25", "2026-2027_504 10∕29∕26"]
    assert (wb.worksheets[0]["G4"].value, wb.worksheets[1]["G4"].value) == ("Before.", "After.")


def test_each_plans_sheet_has_that_plans_date_and_goals(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile().with_plans([OLD_PLAN, NEW_PLAN]))
    store.add_record(student.id, date(2026, 10, 28))
    store.add_record(student.id, date(2026, 10, 29))
    old, new = build_workbook(store, store.records_for_student(student.id), today=TODAY).worksheets
    assert "\nIEP Date: 10/29/2025\n" in old["A1"].value
    assert (old["A2"].value, old["N3"].value) == ("Goal 1: Cut", None)
    assert "\n504 Date: 10/29/2026\n" in new["A1"].value
    assert (new["A2"].value, new["N3"].value) == ("Goal 1: Button\nGoal 2: Write", "Goal 2")


def test_a_new_school_year_in_the_period_starts_a_new_sheet(store):
    student = store.add_student("王小明")
    store.set_profile(student.id, StudentProfile(school_year="09/2026-06/2027", plan_date="10/29/2026"))
    store.add_record(student.id, date(2027, 6, 30), en_text="Summer.")
    store.add_record(student.id, date(2027, 9, 6), en_text="Fall.")
    wb = build_workbook(store, store.records_for_student(student.id), today=TODAY)
    assert wb.sheetnames == ["2026-2027_IEP 10∕29∕26", "2027-2028_IEP 10∕29∕26"]
    assert wb.worksheets[1]["A1"].value.endswith("School year: 09/2027-06/2028")
